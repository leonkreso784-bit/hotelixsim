# -*- coding: utf-8 -*-
"""Excel pregled banke ljudi: listovi Radnici, Gosti, Vanjski. Pokretanje (PowerShell, ne Git Bash):
  $env:PYTHONIOENCODING='utf-8'; py -3 scripts/ljudi-excel.py [korijen-ljudi] [izlaz.xlsx]
Zadano: content/hotel/v2/ljudi i .superpowers/sdd/banka-ljudi-k1/banka-ljudi.xlsx (situation_templates.json i names.json
se traže u roditeljskoj mapi korijena)."""
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

root = os.path.abspath(sys.argv[1] if len(sys.argv) > 1 else "content/hotel/v2/ljudi")
out = os.path.abspath(sys.argv[2] if len(sys.argv) > 2 else ".superpowers/sdd/banka-ljudi-k1/banka-ljudi.xlsx")
v2 = os.path.dirname(root)


def load(path):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def serije(vrsta):
    d = os.path.join(root, vrsta)
    for fn in sorted(os.listdir(d)):
        if fn.startswith("serija-") and fn.endswith(".json"):
            s = load(os.path.join(d, fn))
            for z in s["zapisi"]:
                yield s["serija"], z


def jez(z):
    return ", ".join("%s %s" % (k, r) for k, r in z.get("jezici", []))


def sheet(wb, naslov, zaglavlje, redci, opis_idx, first=False, uvod=None):
    ws = wb.active if first else wb.create_sheet()
    ws.title = naslov
    start = 1
    if uvod:
        ws.cell(row=1, column=1, value=uvod).font = Font(bold=True)
        start = 3
    for c, h in enumerate(zaglavlje, 1):
        cell = ws.cell(row=start, column=c, value=h)
        cell.font = Font(bold=True)
    for r, red in enumerate(redci, start + 1):
        for c, v in enumerate(red, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.alignment = Alignment(vertical="top", wrap_text=(c - 1) in opis_idx)
    for c in range(1, len(zaglavlje) + 1):
        maks = max([len(str(zaglavlje[c - 1]))] + [len(str(r[c - 1] or "")) for r in redci])
        ws.column_dimensions[get_column_letter(c)].width = 70 if (c - 1) in opis_idx else min(max(maks + 2, 8), 40)
    ws.freeze_panes = ws.cell(row=start + 1, column=1)
    ws.auto_filter.ref = "A%d:%s%d" % (start, get_column_letter(len(zaglavlje)), start + max(len(redci), 1))
    return ws


wb = Workbook()
rad = [
    [z["id"], s, z["ime"], z["prezime"], z["rod"], z["dob"], z["uloga"], z["razina"], z["godine_iskustva"],
     z["podrijetlo"], z["zemlja"], z["mjesto"], jez(z), z["opis"], "", ""]
    for s, z in serije("radnici")
]
sheet(wb, "Radnici",
      ["id", "serija", "ime", "prezime", "rod", "dob", "uloga", "razina", "godine iskustva", "podrijetlo", "zemlja",
       "mjesto", "jezici", "opis", "Odluka", "Napomena"], rad, {13}, first=True)

gos = [
    [z["id"], s, z["ime"], z["prezime"], z["rod"], z["dob"], z["zemlja"], z["mjesto"], z.get("segment", ""),
     z.get("zanimanje", ""), jez(z), z["opis"], "", ""]
    for s, z in serije("gosti")
]
sheet(wb, "Gosti",
      ["id", "serija", "ime", "prezime", "rod", "dob", "zemlja", "mjesto", "segment", "zanimanje", "jezici", "opis",
       "Odluka", "Napomena"], gos, {11})

pred = load(os.path.join(v2, "situation_templates.json"))["templates"]
imena = load(os.path.join(v2, "names.json"))["partner"]
van = []
for t in pred:
    for kljuc, r in t.get("roles", {}).items():
        if r.get("from") == "partner":
            opis = r.get("title") or "; ".join(r.get("types", []))
            van.append(["uloga", t["code"], kljuc, r.get("gender", ""), opis])
for rod in ("f", "m"):
    for ime in imena.get(rod, []):
        van.append(["ime", "", "", rod, ime])
sheet(wb, "Vanjski", ["vrsta retka", "predložak", "ključ uloge", "rod", "opis / ime"], van, {4},
      uvod="Vrsta 'vanjski' u banci se tek gradi (nacrt u izradi); ovo su današnje uloge i imena iz igre")

os.makedirs(os.path.dirname(out), exist_ok=True)
wb.save(out)

from openpyxl import load_workbook

chk = load_workbook(out)
for ws in chk.worksheets:
    print("%s: redaka %d (s zaglavljem)" % (ws.title, ws.max_row))
print("xlsx: " + out)
